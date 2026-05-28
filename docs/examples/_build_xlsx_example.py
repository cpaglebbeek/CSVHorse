"""Bouw multi-table-example.xlsx met native Excel cell-styles + _relations + _query.

Pendant van docs/examples/multi-table-example.csv, maar dan met:
- Native Excel bold/color/background per cel (geen __style_*-kolommen)
- _relations-sheet (FK-tabel)
- _query-sheet (laatste query)

Bedoeld als test-case voor SheetHorse XLSX-import (v0.6.0.1-Hanoverian).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

wb = Workbook()
wb.remove(wb.active)

# --- Helpers ---
def vip_font():
    return Font(bold=True, color="1A7F37")
def vip_fill():
    return PatternFill(start_color="DAFBE1", end_color="DAFBE1", fill_type="solid")
def naam_bold():
    return Font(bold=True)
def voorraad_low():
    return Font(bold=True, color="CF222E")
def voorraad_low_fill():
    return PatternFill(start_color="FFEBE9", end_color="FFEBE9", fill_type="solid")
def voorraad_zero():
    return Font(bold=True, color="FFFFFF")
def voorraad_zero_fill():
    return PatternFill(start_color="CF222E", end_color="CF222E", fill_type="solid")
def big_order():
    return Font(bold=True, color="0969DA")

# --- Sheet 1: klanten ---
klanten = wb.create_sheet("klanten")
klanten.append(["id", "naam", "stad", "vip"])
rows_k = [
    (1, "Alice Janssen", "Amsterdam", "true"),
    (2, "Bob de Vries", "Rotterdam", "false"),
    (3, "Carla Smit", "Den Haag", "true"),
    (4, "Daan Bakker", "Utrecht", "false"),
    (5, "Esther Visser", "Eindhoven", "true"),
]
for r in rows_k:
    klanten.append(r)
# Style VIP-rijen (rows 2, 4, 6 in Excel-numbering)
for excel_row in (2, 4, 6):  # rij-numbering 1-based, header op rij 1
    klanten.cell(row=excel_row, column=2).font = naam_bold()  # naam
    klanten.cell(row=excel_row, column=4).font = vip_font()   # vip
    klanten.cell(row=excel_row, column=4).fill = vip_fill()

# --- Sheet 2: producten ---
producten = wb.create_sheet("producten")
producten.append(["sku", "naam", "prijs", "voorraad"])
rows_p = [
    ("P001", "Stoel eik", 189.00, 12),
    ("P002", "Tafel notenhout", 449.50, 3),
    ("P003", "Lamp messing", 79.95, 28),
    ("P004", "Kast wit", 329.00, 0),
    ("P005", "Bureau zwart", 259.00, 7),
]
for r in rows_p:
    producten.append(r)
# Voorraad 3 = low (rij 3)
producten.cell(row=3, column=4).font = voorraad_low()
producten.cell(row=3, column=4).fill = voorraad_low_fill()
# Voorraad 0 = zero (rij 5)
producten.cell(row=5, column=4).font = voorraad_zero()
producten.cell(row=5, column=4).fill = voorraad_zero_fill()

# --- Sheet 3: bestellingen ---
bestellingen = wb.create_sheet("bestellingen")
bestellingen.append(["order_id", "klant_id", "sku", "aantal", "datum"])
rows_b = [
    (1001, 1, "P001", 2, "2026-04-12"),
    (1002, 1, "P003", 1, "2026-04-12"),
    (1003, 2, "P002", 1, "2026-05-03"),
    (1004, 3, "P001", 4, "2026-05-08"),
    (1005, 3, "P005", 1, "2026-05-08"),
    (1006, 5, "P003", 3, "2026-05-15"),
    (1007, 5, "P002", 1, "2026-05-15"),
    (1008, 4, "P004", 1, "2026-05-20"),
]
for r in rows_b:
    bestellingen.append(r)
# Grote orders (aantal > 2): rij 5 (1004, 4 stuks) en rij 7 (1006, 3 stuks)
bestellingen.cell(row=5, column=4).font = big_order()
bestellingen.cell(row=7, column=4).font = big_order()

# --- Sheet 4: _relations ---
relations = wb.create_sheet("_relations")
relations.append(["from_sheet", "from_col", "to_sheet", "to_col", "cardinality", "name"])
relations.append(["bestellingen", "klant_id", "klanten", "id", "n-1", "fk_bestelling_klant"])
relations.append(["bestellingen", "sku", "producten", "sku", "n-1", "fk_bestelling_product"])

# --- Sheet 5: _query ---
query = wb.create_sheet("_query")
query.append(["query", "executed_at", "row_count"])
query.append([
    "SELECT k.naam AS klant, k.stad, COUNT(*) AS aantal_orders, "
    "SUM(b.aantal * p.prijs) AS totaal_eur "
    "FROM [klanten] k JOIN [bestellingen] b ON b.klant_id=k.id "
    "JOIN [producten] p ON p.sku=b.sku "
    "WHERE k.vip='true' GROUP BY k.naam, k.stad ORDER BY totaal_eur DESC",
    "2026-05-28T17:00:00.000Z",
    3
])

wb.save("docs/examples/multi-table-example.xlsx")
print("Wrote: docs/examples/multi-table-example.xlsx")
print(f"  Sheets: {wb.sheetnames}")
